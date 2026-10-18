
import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document


class ErrorExtraccionCV(Exception):
    """Error al procesar un currículum."""


def normalizar_texto_pdf(texto):
    """
    Corrige fragmentos de palabras separados artificialmente
    durante la extracción de texto de archivos PDF.
    """
    lineas = []

    correcciones = (
        (r"\bSu\s+bgere\s+nte\b", "Subgerente"),
        (r"\bProye\s+ctos\b", "Proyectos"),
        (r"\bLici\s+ta\s+ci\s+o\s+nes\b", "Licitaciones"),
    )

    for linea in texto.splitlines():
        partes = re.split(r"(\s{2,})", linea)
        partes_normalizadas = []

        for parte in partes:
            if not parte or parte.isspace():
                partes_normalizadas.append(parte)
                continue

            patron = r"(?<!\S)(?:[^\W\d_]\s+){3,}[^\W\d_](?!\S)"

            parte = re.sub(
                patron,
                lambda coincidencia: re.sub(
                    r"\s+", "", coincidencia.group(0)
                ),
                parte,
                flags=re.UNICODE,
            )

            for patron_correccion, reemplazo in correcciones:
                parte = re.sub(
                    patron_correccion,
                    reemplazo,
                    parte,
                    flags=re.IGNORECASE,
                )

            partes_normalizadas.append(parte)

        lineas.append("".join(partes_normalizadas))

    return "\n".join(lineas)


def extraer_texto_pdf(archivo):
    """Extrae texto de PDF intentando conservar su disposición visual."""
    lector = PdfReader(archivo, strict=False)
    paginas = []

    for pagina in lector.pages:
        # Las páginas en blanco pueden omitir /Contents en un PDF válido.
        if pagina.get_contents() is None:
            continue
        try:
            texto = pagina.extract_text(extraction_mode="layout") or ""
        except (ValueError, TypeError, NotImplementedError):
            texto = ""

        if not texto.strip():
            texto = pagina.extract_text() or ""

        if texto.strip():
            paginas.append(normalizar_texto_pdf(texto.strip()))

    return "\n\n".join(paginas)


def extraer_texto_docx(archivo):
    """Extrae párrafos y tablas de un documento Word."""
    documento = Document(archivo)
    textos = []

    for parrafo in documento.paragraphs:
        if parrafo.text.strip():
            textos.append(parrafo.text.strip())

    for tabla in documento.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda.text.strip():
                    textos.append(celda.text.strip())

    return "\n\n".join(textos)


def extraer_texto_cv(archivo, nombre_archivo):
    """Selecciona el lector según la extensión del archivo."""
    extension = Path(nombre_archivo).suffix.lower()

    if extension not in (".pdf", ".docx"):
        raise ErrorExtraccionCV(
            "Formato no compatible con la extracción automática."
        )

    try:
        if extension == ".pdf":
            texto = extraer_texto_pdf(archivo)
        else:
            texto = extraer_texto_docx(archivo)

    except Exception as error:
        raise ErrorExtraccionCV(
            "No fue posible extraer el texto del currículum."
        ) from error

    # PostgreSQL no admite NUL en campos text; el archivo original se conserva.
    return texto.replace("\x00", "")
